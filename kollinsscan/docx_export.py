"""Word (.docx) export: title page, a clickable table of contents, heading
styles, and each chapter on a new page."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .document import book_blocks, heading_title

FONT = "Times New Roman"


def _el(tag: str, **attrs) -> OxmlElement:
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), str(value))
    return el


def _add_runs(paragraph, runs: list[dict]) -> None:
    for r in runs:
        for n, piece in enumerate(r["text"].split("\n")):
            if n:
                paragraph.add_run().add_break(WD_BREAK.LINE)
            if piece:
                run = paragraph.add_run(piece)
                run.bold = r["b"] or None
                run.italic = r["i"] or None
                run.underline = r["u"] or None


def _bookmark(paragraph, name: str, bookmark_id: int) -> None:
    p = paragraph._p
    start = _el("w:bookmarkStart", id=bookmark_id, name=name)
    end = _el("w:bookmarkEnd", id=bookmark_id)
    p.insert(1 if p.pPr is not None else 0, start)
    p.append(end)


def _toc_entry(paragraph, title: str, page: str, anchor: str) -> None:
    """One line of the table of contents, linked to its heading."""
    link = _el("w:hyperlink", anchor=anchor, history=1)
    for text, is_tab in ((title, False), (None, True), (page, False)):
        r = _el("w:r")
        if is_tab:
            r.append(_el("w:tab"))
        else:
            t = _el("w:t")
            t.text = text
            t.set(qn("xml:space"), "preserve")
            r.append(t)
        link.append(r)
    paragraph._p.append(link)


def _field_char(paragraph, kind: str) -> None:
    r = _el("w:r")
    r.append(_el("w:fldChar", fldCharType=kind))
    paragraph._p.append(r)


def _instr(paragraph, text: str) -> None:
    r = _el("w:r")
    t = _el("w:instrText")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    paragraph._p.append(r)


def to_docx(book: dict, pages: list[dict]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5
    for name, size, centered in (("Heading 1", 18, True), ("Heading 2", 15, True),
                                 ("Heading 3", 13, False)):
        style = styles[name]
        style.font.name = FONT
        fonts = style.element.rPr.rFonts
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)  # theme fonts override the name
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(24 if name == "Heading 1" else 12)
        style.paragraph_format.space_after = Pt(12)
        if centered:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 1"].paragraph_format.page_break_before = True

    doc.core_properties.title = book["title"]
    doc.core_properties.author = book.get("author", "")

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(144)
    run = title.add_run(book["title"])
    run.bold, run.font.size = True, Pt(26)
    if book.get("author"):
        author = doc.add_paragraph()
        author.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author.add_run(book["author"]).font.size = Pt(16)

    blocks = book_blocks(pages)
    headings = [(n, b, label) for n, (b, label) in enumerate(blocks) if b["tag"] != "p"]

    # Table of contents: a real TOC field (Word refreshes it with this file's
    # own page numbers when the file opens), pre-filled with the printed
    # book's page numbers and links to each heading.
    if headings:
        caption = doc.add_paragraph()
        caption.paragraph_format.page_break_before = True
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run("Contents")
        run.bold, run.font.size = True, Pt(16)
        for k, (n, block, label) in enumerate(headings):
            entry = doc.add_paragraph()
            fmt = entry.paragraph_format
            fmt.left_indent = Inches(0.25 * (int(block["tag"][1]) - 1))
            fmt.space_after = Pt(2)
            fmt.line_spacing = 1.15
            fmt.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if k == 0:
                _field_char(entry, "begin")
                _instr(entry, ' TOC \\o "1-3" \\h \\z \\u ')
                _field_char(entry, "separate")
            _toc_entry(entry, heading_title(block), label, f"_TocKS{n}")
            if k == len(headings) - 1:
                _field_char(entry, "end")
        doc.settings.element.append(_el("w:updateFields", val="true"))
    # Chapter headings start a new page by themselves; anything else needs
    # a page break after the title or contents page.
    if not blocks or blocks[0][0]["tag"] != "h1":
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for n, (block, _) in enumerate(blocks):
        if block["tag"] == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.3)
        else:
            p = doc.add_paragraph(style=f"Heading {block['tag'][1]}")
            _bookmark(p, f"_TocKS{n}", n + 1)
        _add_runs(p, block["runs"])

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
